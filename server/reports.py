"""报告渲染：自包含 HTML 报告 + Word(.docx) 报告。"""
from __future__ import annotations

import html as _html
from datetime import datetime

STATUS_META = {
    "normal": ("正常", "#047857", "#e8f8f1"),
    "warn": ("警告", "#b45309", "#fef6e7"),
    "crit": ("严重", "#b91c1c", "#fdecec"),
    "info": ("信息", "#4338ca", "#eef2ff"),
    "error": ("错误", "#b91c1c", "#fdecec"),
}


def _esc(value) -> str:
    return _html.escape("" if value is None else str(value))


def _status_badge(status: str) -> str:
    label, color, bg = STATUS_META.get(status, (status, "#475569", "#f1f5f9"))
    return (f'<span style="display:inline-block;padding:2px 10px;border-radius:999px;'
            f'font-size:12px;font-weight:600;color:{color};background:{bg}">{_esc(label)}</span>')


def _summary_cards(summary: dict) -> str:
    cards = [
        ("巡检项总数", summary.get("total", 0), "#111827"),
        ("正常", summary.get("normal", 0), "#059669"),
        ("警告", summary.get("warn", 0), "#d97706"),
        ("严重", summary.get("crit", 0), "#dc2626"),
        ("信息", summary.get("info", 0), "#4338ca"),
        ("错误", summary.get("error", 0), "#dc2626"),
    ]
    out = []
    for label, num, color in cards:
        out.append(
            f'<div style="flex:1;border:1px solid #e5e7eb;border-radius:10px;padding:14px;'
            f'text-align:center;background:#fafbfc"><div style="font-size:24px;font-weight:700;'
            f'color:{color};line-height:1.3">{num}</div>'
            f'<div style="font-size:12px;color:#6b7280;margin-top:2px">{label}</div></div>'
        )
    return f'<div id="summary" style="display:flex;gap:12px;margin:0 0 26px">{ "".join(out) }</div>'


def _abnormal_summary(result: dict) -> str:
    """巡检结果汇总：异常明细 + 失败信息。"""
    abnormal = []
    failures = []
    for cat in result.get("categories", []):
        for item in cat.get("items", []):
            st = item.get("status")
            if st == "error":
                failures.append((cat.get("name"), item))
            elif st in ("warn", "crit"):
                abnormal.append((cat.get("name"), item))

    th = ('style="text-align:left;font-weight:600;color:#334155;font-size:12px;'
          'padding:8px 10px;border-bottom:1px solid #e5e7eb;background:#f1f5f9"')
    td = 'style="padding:7px 10px;border-bottom:1px solid #f1f5f9;vertical-align:top"'
    out = []

    crit_n = sum(1 for _, i in abnormal if i.get("status") == "crit")
    warn_n = sum(1 for _, i in abnormal if i.get("status") == "warn")
    total_abn = len(abnormal) + len(failures)
    if total_abn:
        out.append(
            f'<div style="margin:20px 0 6px;padding:12px 16px;background:#fef2f2;border:1px solid #fecaca;'
            f'border-radius:10px;font-size:13.5px;font-weight:600;color:#b91c1c">'
            f'⚠ 巡检项异常汇总：共 {total_abn} 项（严重 {crit_n} · 警告 {warn_n} · 失败 {len(failures)}）</div>'
        )

    if abnormal:
        rows = "".join(
            f'<tr><td {td}>{_esc(c)}</td><td {td}>{_esc(i.get("name"))}</td>'
            f'<td {td}>{_esc(i.get("value"))}</td>'
            f'<td {td} style="padding:7px 10px;border-bottom:1px solid #f1f5f9;color:{STATUS_META.get(i.get("status"), ("", "#475569"))[1]};font-weight:600">{_esc(STATUS_META.get(i.get("status"), (i.get("status"),))[0])}</td>'
            f'<td {td} style="color:#6b7280">{_esc(i.get("threshold_desc") or "—")}</td></tr>'
            for c, i in abnormal
        )
        out.append(
            f'<h2 style="font-size:15px;font-weight:600;margin:26px 0 10px;padding-left:10px;'
            f'border-left:3px solid #2563eb">巡检结果汇总 · 异常明细</h2>'
            f'<table style="width:100%;border-collapse:collapse">'
            f'<thead><tr><th {th}>分类</th><th {th}>巡检项</th><th {th}>巡检结果</th>'
            f'<th {th}>状态</th><th {th}>阈值规则</th></tr></thead><tbody>{rows}</tbody></table>'
        )

    if failures:
        rows = "".join(
            f'<tr><td {td}>{_esc(c)}</td><td {td}>{_esc(i.get("name"))}</td>'
            f'<td {td} style="color:#b91c1c">{_esc(i.get("error") or i.get("value"))}</td></tr>'
            for c, i in failures
        )
        out.append(
            f'<h2 style="font-size:15px;font-weight:600;margin:26px 0 10px;padding-left:10px;'
            f'border-left:3px solid #dc2626">失败信息</h2>'
            f'<table style="width:100%;border-collapse:collapse">'
            f'<thead><tr><th {th}>分类</th><th {th}>巡检项</th><th {th}>失败原因</th></tr></thead>'
            f'<tbody>{rows}</tbody></table>'
        )

    return "".join(out)


def _item_table(item: dict, abnormal_only: bool = False) -> str:
    columns = item.get("columns") or []
    rows = item.get("rows") or []
    if columns and rows:
        thead = "".join(
            f'<th style="text-align:left;font-weight:600;color:#334155;font-size:12px;'
            f'padding:8px 10px;border-bottom:1px solid #e5e7eb;background:#f1f5f9">{_esc(c)}</th>'
            for c in columns[:8]
        )
        body = ""
        for r in rows[:50]:
            tds = "".join(
                f'<td style="padding:7px 10px;border-bottom:1px solid #f1f5f9;font-size:12.5px;'
                f'word-break:break-all">{_esc(r[i]) if i < len(r) else ""}</td>'
                for i in range(len(columns[:8]))
            )
            body += f"<tr>{tds}</tr>"
        table = (f'<table style="width:100%;border-collapse:collapse;margin-top:6px">'
                 f'<thead><tr>{thead}</tr></thead><tbody>{body}</tbody></table>')
        return table
    return ""


def render_html(result: dict, report_no: str, trigger_label: str = "手动巡检",
                abnormal_only: bool = False) -> str:
    conn = result.get("connection", {})
    tpl = result.get("template", {})
    summary = result.get("summary", {})
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    categories_html = []
    for idx, cat in enumerate(result.get("categories", [])):
        items_html = []
        for item in cat.get("items", []):
            if abnormal_only and item.get("status") in ("normal", "info"):
                continue
            items_html.append(
                f'<h3 style="font-size:13px;font-weight:600;color:#111827;margin:14px 0 6px">'
                f'{_esc(item.get("name"))} '
                f'<span style="font-weight:400;font-size:12px;color:#6b7280">'
                f'（{_esc(item.get("threshold_desc") or "仅采集展示")}）</span></h3>'
                f'<div style="margin-bottom:2px">{_status_badge(item.get("status"))} '
                f'<span style="font-size:13px;color:#111827">{_esc(item.get("value"))}</span></div>'
                + _item_table(item)
            )
        if items_html:
            categories_html.append(
                f'<h2 id="cat-{idx}" style="font-size:15px;font-weight:600;margin:26px 0 10px;padding-left:10px;'
                f'border-left:3px solid #2563eb">{_esc(cat.get("name"))}</h2>'
                + "".join(items_html)
            )

    if not categories_html:
        categories_html.append('<p style="color:#6b7280">无巡检数据</p>')

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{_esc(tpl.get("report_title") or "数据库巡检报告")}</title>
<style>
  body{{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Microsoft YaHei",sans-serif;
       color:#111827;font-size:14px;line-height:1.6;margin:0;padding:0;background:#fff}}
  .page{{max-width:1000px;margin:0 auto;padding:34px 40px}}
  .header{{display:flex;justify-content:space-between;align-items:flex-start;border-bottom:2px solid #111827;padding-bottom:18px;margin-bottom:22px}}
  .header h1{{font-size:22px;margin:0 0 6px}}
  .header p{{font-size:12.5px;color:#6b7280;margin:2px 0}}
  .meta{{text-align:right;font-size:12px;color:#6b7280;line-height:1.9}}
  .footer{{margin-top:32px;padding-top:14px;border-top:1px dashed #e5e7eb;font-size:11.5px;color:#6b7280;display:flex;justify-content:space-between}}
  @media print{{.page{{max-width:none;padding:20px}}}}
</style>
</head>
<body>
<div class="page">
  <div class="header">
    <div>
      <h1>{_esc(tpl.get("report_title") or "数据库巡检报告")}</h1>
      <p>实例：{_esc(conn.get("name"))}（{_esc(conn.get("version") or conn.get("db_type") or "")}）</p>
      <p>巡检时间：{_esc(result.get("started_at"))} &nbsp;|&nbsp; 耗时：{_esc(result.get("duration"))} 秒</p>
    </div>
    <div class="meta">
      报告编号：{_esc(report_no)}<br>
      巡检模板：{_esc(tpl.get("name"))}<br>
      执行方式：{_esc(trigger_label)}<br>
      生成时间：{_esc(now_str)}
    </div>
  </div>
  {_summary_cards(summary)}
  {_abnormal_summary(result)}
  { "".join(categories_html) }
  <div class="footer">
    <span>本报告由 Bond-DBCheck 邦德智能巡检平台自动生成 · {_esc(tpl.get("watermark") or "")}</span>
    <span>{_esc(now_str)}</span>
  </div>
</div>
</body>
</html>"""


# ============================= Word =============================
def render_word(result: dict, report_no: str, trigger_label: str = "手动巡检",
                abnormal_only: bool = False) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    conn = result.get("connection", {})
    tpl = result.get("template", {})
    summary = result.get("summary", {})

    def cn_font(run, size=None, bold=None, color=None, name="微软雅黑"):
        run.font.name = name
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), name)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if color is not None:
            run.font.color.rgb = color

    def shade_cell(cell, fill="1F4E79"):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def style_table(table):
        """表格自适应宽度并居中，协调美观。"""
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(9)

    # ---- 标题与元信息 ----
    title = doc.add_heading(tpl.get("report_title") or "数据库巡检报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in title.runs:
        cn_font(r, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = meta.add_run(
        f"实例：{conn.get('name')}（{conn.get('version') or conn.get('db_type') or ''}）\n"
        f"巡检时间：{result.get('started_at')}    耗时：{result.get('duration')} 秒\n"
        f"报告编号：{report_no}    巡检模板：{tpl.get('name')}    执行方式：{trigger_label}"
    )
    cn_font(m, size=8, color=RGBColor(0x6B, 0x72, 0x80))

    # ---- 巡检汇总 ----
    h = doc.add_heading("巡检汇总", level=1)
    for r in h.runs:
        cn_font(r, size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    p = doc.add_paragraph()
    run = p.add_run(
        f"巡检项总数 {summary.get('total', 0)} ｜ 正常 {summary.get('normal', 0)} ｜ "
        f"警告 {summary.get('warn', 0)} ｜ 严重 {summary.get('crit', 0)} ｜ "
        f"信息 {summary.get('info', 0)} ｜ 错误 {summary.get('error', 0)}"
    )
    cn_font(run, size=10, bold=True)

    # 异常明细
    abnormal = [(c.get("name"), it) for c in result.get("categories", [])
                for it in c.get("items", []) if it.get("status") in ("warn", "crit")]
    failures = [(c.get("name"), it) for c in result.get("categories", [])
                for it in c.get("items", []) if it.get("status") == "error"]

    crit_n = sum(1 for _, it in abnormal if it.get("status") == "crit")
    warn_n = sum(1 for _, it in abnormal if it.get("status") == "warn")
    total_abn = len(abnormal) + len(failures)
    if total_abn:
        p = doc.add_paragraph()
        run = p.add_run(f"⚠ 巡检项异常汇总：共 {total_abn} 项（严重 {crit_n} · 警告 {warn_n} · 失败 {len(failures)}）")
        cn_font(run, size=11, bold=True, color=RGBColor(0xB9, 0x1C, 0x1C))

    if abnormal:
        p = doc.add_paragraph()
        run = p.add_run("异常明细")
        cn_font(run, size=11, bold=True, color=RGBColor(0xB4, 0x53, 0x09))
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for i, htext in enumerate(["分类", "巡检项", "巡检结果", "阈值规则", "状态"]):
            hdr[i].text = htext
            for rr in hdr[i].paragraphs[0].runs:
                cn_font(rr, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade_cell(hdr[i], "D97706")
        for cname, it in abnormal:
            row = tbl.add_row().cells
            row[0].text = cname
            row[1].text = it.get("name", "")
            row[2].text = str(it.get("value") or "")
            row[3].text = it.get("threshold_desc") or "—"
            label, color_hex, _ = STATUS_META.get(it.get("status"), (it.get("status"), "#475569", ""))
            row[4].text = label
            for rr in row[4].paragraphs[0].runs:
                cn_font(rr, size=9, bold=True, color=RGBColor.from_string(color_hex[1:]))
        style_table(tbl)

    if failures:
        p = doc.add_paragraph()
        run = p.add_run("失败信息")
        cn_font(run, size=11, bold=True, color=RGBColor(0xB9, 0x1C, 0x1C))
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for i, htext in enumerate(["分类", "巡检项", "失败原因"]):
            hdr[i].text = htext
            for rr in hdr[i].paragraphs[0].runs:
                cn_font(rr, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade_cell(hdr[i], "B91C1C")
        for cname, it in failures:
            row = tbl.add_row().cells
            row[0].text = cname
            row[1].text = it.get("name", "")
            row[2].text = str(it.get("error") or it.get("value") or "")
        style_table(tbl)

    # ---- 分类明细 ----
    cn_num = "一二三四五六七八九十"
    for idx, cat in enumerate(result.get("categories", [])):
        prefix = cn_num[idx] if idx < len(cn_num) else str(idx + 1)
        hcat = doc.add_heading(f"{prefix}、{cat.get('name')}", level=1)
        for r in hcat.runs:
            cn_font(r, size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

        items = [it for it in cat.get("items", [])
                 if not (abnormal_only and it.get("status") in ("normal", "info"))]
        if not items:
            continue
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, htext in enumerate(["巡检项", "巡检结果", "阈值", "状态"]):
            hdr[i].text = htext
            for rr in hdr[i].paragraphs[0].runs:
                cn_font(rr, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade_cell(hdr[i], "1F4E79")
        for item in items:
            row = table.add_row().cells
            row[0].text = item.get("name", "")
            detail = item.get("rows") or []
            cols = item.get("columns") or []
            value_text = str(item.get("value") or "")
            if detail and cols:
                lines = []
                for r in detail[:5]:
                    parts = [f"{cols[i]}={r[i]}" for i in range(min(len(cols), len(r), 6))]
                    lines.append(" | ".join(parts))
                if len(detail) > 5:
                    lines.append(f"… 共 {len(detail)} 行")
                value_text += "\n" + "\n".join(lines)
            row[1].text = value_text
            row[2].text = item.get("threshold_desc") or "—"
            label, color_hex, _ = STATUS_META.get(item.get("status"), (item.get("status"), "#475569", ""))
            row[3].text = label
            for rr in row[3].paragraphs[0].runs:
                cn_font(rr, size=9, bold=True, color=RGBColor.from_string(color_hex[1:]))
        style_table(table)

    footer = doc.add_paragraph()
    fr = footer.add_run(
        f"本报告由 Bond-DBCheck 邦德智能巡检平台自动生成 · {tpl.get('watermark') or ''}  ·  "
        f"生成时间 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    cn_font(fr, size=8, color=RGBColor(0x6B, 0x72, 0x80))

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
